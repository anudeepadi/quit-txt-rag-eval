"""Run real split/statistics helpers on wholly invented fixture data.

This is an onboarding check, not a RAGAS/model run or a study result. No API
keys, study files, model downloads, or network connections are needed at run time.
"""

import importlib.util
import json
import os
from pathlib import Path
import socket
import sys
import tempfile
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "code"))


def deny_network(*args, **kwargs):
    raise RuntimeError("Network access is disabled for this synthetic example")


def main():
    from openpyxl import Workbook
    from shared.question_split import (
        load_dev_questions, load_test_questions_frozen, TestSetLeakError,
    )
    from shared.ragas_utils import compute_faithfulness_stats

    spec = importlib.util.spec_from_file_location(
        "holdout", ROOT / "code/scripts/holdout_noninferiority.py"
    )
    holdout = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(holdout)

    # Use fabricated nonmedical questions. Match the real loader's workbook
    # contract: question A, reference B, reviewer exclusion flag CE (column 83).
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "synthetic-questions.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(1, 1, "Question")
        sheet.cell(1, 2, "Reference")
        sheet.cell(1, 83, "Comments LFV")
        for index in range(61):
            row = index + 2
            sheet.cell(row, 1, f"Synthetic fixture {index}: what is the label?")
            sheet.cell(row, 2, f"Invented label {index}.")
        sheet.cell(62, 83, "Synthetic exclusion flag")
        workbook.save(path)
        workbook.close()

        # Explicitly clear any inherited optimization marker only for this
        # fabricated workbook. The real frozen study data is never accessed.
        with patch.dict(os.environ, {"EVO_OPTIMIZING": "0"}):
            dev = load_dev_questions(path)
            test = load_test_questions_frozen(path)
        assert (len(dev), len(test)) == (50, 10)
        assert not ({q.question for q in dev} & {q.question for q in test})
        with patch.dict(os.environ, {"EVO_OPTIMIZING": "1"}):
            try:
                load_test_questions_frozen(path)
            except TestSetLeakError:
                guard_passed = True
            else:
                raise AssertionError("Frozen-test access must be blocked during optimization")

    fixture = json.loads((ROOT / "examples/synthetic_scores.json").read_text())
    ai = fixture["ai_rag"]
    human = fixture["human_rag"]
    assert len(ai) == len(human) == len(test)
    assert all(0 <= value <= 1 for value in ai + human)
    differences = [a - h for a, h in zip(ai, human)]
    ni = holdout.paired_noninferiority(differences, fixture["illustrative_margin"])
    tost = holdout.paired_tost(differences, fixture["illustrative_margin"])

    print("# Synthetic onboarding result — NOT study findings\n")
    print("These scores are hand-authored fixtures, not generated or judged answers.")
    print("The run exercises the repository's actual loader, split guard, summary, and paired tests.\n")
    print("Split: 61 invented rows → 1 flagged row removed → 50 dev / 10 held out.")
    print(f"Optimization guard: {'PASS' if guard_passed else 'FAIL'}. Dev/test overlap: 0.\n")
    print("| Fixture | n | Mean | SD | Descriptive 95% interval |")
    print("| --- | ---: | ---: | ---: | --- |")
    for label, scores in (("AI RAG fixture", ai), ("Human RAG fixture", human)):
        result = compute_faithfulness_stats(scores)
        print(f"| {label} | {result['n']} | {result['mean']:.4f} | {result['sd']:.4f} | "
              f"{result['ci_lower']:.4f}–{result['ci_upper']:.4f} |")
    print("\nIntervals above use the existing helper's normal approximation; n=10 is only a fixture.")
    print(f"\nPaired mean difference (AI − human): {ni['mean_diff']:+.4f}.")
    print(f"Illustrative margin: {fixture['illustrative_margin']:.2f}; "
          f"non-inferiority p={ni['p_value']:.6f}; TOST p={tost['p_value']:.6f}.")
    print("These calculations do not establish model parity or clinical effectiveness.")
    print("The study's 50/77 split and recorded manifest are not changed by this example.")


if __name__ == "__main__":
    with patch.object(socket.socket, "connect", deny_network), patch.object(socket.socket, "connect_ex", deny_network):
        main()
